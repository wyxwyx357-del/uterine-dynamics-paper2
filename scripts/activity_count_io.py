"""Preserve raw clinician activity-count Excel cells before pandas NA parsing."""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ACTIVITY_COLUMNS = ("peristalsis_forward_count", "peristalsis_reverse_count")


def read_excel_preserving_activity(path: Path, sheet_name: str, **kwargs) -> pd.DataFrame:
    """Parse other fields normally; restore the two original count cell values.

    openpyxl also preserves native Excel error cells such as #N/A, which
    pandas converts to NaN even when keep_default_na=False.
    """
    table = pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl", **kwargs)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook[sheet_name]
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        for name in ACTIVITY_COLUMNS:
            indices = [i for i, value in enumerate(headers)
                       if isinstance(value, str) and value.split("\n", 1)[0] == name]
            columns = [column for column in table.columns
                       if str(column).split("\n", 1)[0] == name]
            if not indices and not columns:
                continue
            if len(indices) != 1 or len(columns) != 1:
                raise ValueError(f"{sheet_name}: expected one activity-count column {name}")
            raw_values = [row[indices[0]] for row in sheet.iter_rows(
                min_row=2, max_row=len(table) + 1, values_only=True)]
            table[columns[0]] = pd.Series(raw_values, index=table.index, dtype=object)
    finally:
        workbook.close()
    return table
